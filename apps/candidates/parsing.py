from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from django.core.files.base import File

MAX_SKILLS = 20
MAX_SKILL_LENGTH = 40
MAX_YEARS = 60

_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE = re.compile(r"(?:\+84|0084|0)[\s.-]?(?:\d[\s.-]?){8,10}\d")
_YEARS = re.compile(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)\b[^.\n]{0,20}?experience", re.I)
_TITLE_LABEL = re.compile(r"^(?:current\s+)?(?:title|position|job\s+title)\s*[:\-]\s*(.+)$", re.I)
_SKILLS_HEADING = re.compile(
    r"^(?:technical\s+)?(?:skills|skill\s+set|technologies)\s*[:\-]?\s*(.*)$", re.I
)
_SECTION_HEADING = re.compile(
    r"^(summary|experience|education|projects|languages|certifications|interests|awards)\b", re.I
)
_SPLIT_SKILLS = re.compile(r"[,;|·•/]| {2,}")
_CONTACT_HINT = re.compile(r"@|https?://|linkedin|\d", re.I)


def _pdf_text(upload: File[Any]) -> str:
    from pypdf import PdfReader

    reader = PdfReader(upload)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(upload: File[Any]) -> str:
    from docx import Document

    document = Document(upload)
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def extract_text(file: File[Any]) -> str:
    extension = Path(file.name or "").suffix.lower()
    reader = {".pdf": _pdf_text, ".docx": _docx_text}.get(extension)
    if reader is None:
        return ""
    file.seek(0)
    try:
        return reader(file).strip()
    except Exception:
        return ""
    finally:
        file.seek(0)


def _guess_phone(text: str) -> str:
    match = _PHONE.search(text)
    if not match:
        return ""
    digits = re.sub(r"\D", "", match.group(0))
    if digits.startswith("0084"):
        digits = digits[4:]
    elif digits.startswith("84"):
        digits = digits[2:]
    elif digits.startswith("0"):
        digits = digits[1:]
    return f"+84{digits}"


def _guess_name_and_title(lines: list[str]) -> tuple[str, str]:
    labelled = ""
    for line in lines:
        match = _TITLE_LABEL.match(line)
        if match:
            labelled = match.group(1).strip()[:160]
            break
    for index, line in enumerate(lines):
        words = line.split()
        if not (2 <= len(words) <= 6 and len(line) <= 60):
            continue
        if _CONTACT_HINT.search(line) or _SECTION_HEADING.match(line):
            continue
        if labelled:
            return line, labelled
        for follower in lines[index + 1 : index + 3]:
            if (
                follower
                and not _SECTION_HEADING.match(follower)
                and not _CONTACT_HINT.search(follower)
                and len(follower) <= 160
            ):
                return line, follower
        return line, ""
    return "", labelled


def _guess_skills(lines: list[str]) -> list[str]:
    collected: list[str] = []
    capturing = False
    for line in lines:
        heading = _SKILLS_HEADING.match(line)
        if heading:
            capturing = True
            collected.extend(_SPLIT_SKILLS.split(heading.group(1)))
            continue
        if capturing:
            if not line or _SECTION_HEADING.match(line):
                break
            collected.extend(_SPLIT_SKILLS.split(line))
    skills: list[str] = []
    seen: set[str] = set()
    for raw in collected:
        skill = raw.strip(" .-\t")[:MAX_SKILL_LENGTH]
        if skill and skill.lower() not in seen:
            skills.append(skill)
            seen.add(skill.lower())
    return skills[:MAX_SKILLS]


def guess_fields(text: str) -> dict[str, Any]:
    lines = [line.strip() for line in (text or "").splitlines()]
    full_name, current_title = _guess_name_and_title(lines)
    email = _EMAIL.search(text)
    years = _YEARS.search(text)
    years_value = int(years.group(1)) if years else None
    if years_value is not None and years_value > MAX_YEARS:
        years_value = None
    return {
        "full_name": full_name,
        "current_title": current_title,
        "email": email.group(0).lower() if email else "",
        "phone": _guess_phone(text),
        "skills": _guess_skills(lines),
        "years_experience": years_value,
    }
