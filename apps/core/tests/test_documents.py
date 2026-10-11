import base64
import io
from types import SimpleNamespace
from typing import Any

import anthropic
import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from pydantic import BaseModel

from apps.core.documents import (
    MAX_DOCUMENT_BYTES,
    extract_text,
    is_scanned,
    read_with_claude,
    source_block,
    validate_document,
)
from apps.core.testing.pdf import make_pdf


class Shape(BaseModel):
    title: str | None


def ask(response: Any = None, error: Exception | None = None, **overrides: Any) -> Any:
    def parse(**_: Any) -> Any:
        if error is not None:
            raise error
        return response

    return read_with_claude(
        SimpleNamespace(messages=SimpleNamespace(parse=parse)),
        **{
            "model": "claude-sonnet-5-5",
            "source": {"type": "text", "text": "<jd>x</jd>"},
            "instructions": "Read it.",
            "output_format": Shape,
            "max_tokens": 100,
            "output_config": {"effort": "low"},
            "label": "JD",
            **overrides,
        },
    )


def reply(stop_reason: str = "end_turn", title: str = "Kế toán trưởng") -> SimpleNamespace:
    return SimpleNamespace(
        stop_reason=stop_reason,
        parsed_output=Shape(title=title) if stop_reason == "end_turn" else None,
        model="claude-sonnet-5-5-20261001",
        usage=SimpleNamespace(input_tokens=900, output_tokens=120),
        _request_id="req_doc",
    )


def test_the_caller_names_the_document_in_type_and_size_errors() -> None:
    with pytest.raises(ValidationError, match="as a JD"):
        validate_document(SimpleUploadedFile("jd.png", b"\x89PNG"), "Upload it as a JD.", "big")
    with pytest.raises(ValidationError, match="JDs must be small"):
        validate_document(
            SimpleUploadedFile("jd.pdf", make_pdf("x") + b"0" * MAX_DOCUMENT_BYTES),
            "type",
            "JDs must be small",
        )


def test_a_pdf_text_layer_is_extracted_and_the_pointer_rewound() -> None:
    upload = SimpleUploadedFile("jd.pdf", make_pdf("Senior Accountant\nHo Chi Minh City"))

    assert "Senior Accountant" in extract_text(upload)
    assert upload.tell() == 0


def test_only_a_pdf_without_text_counts_as_scanned() -> None:
    assert is_scanned("jd.pdf", "")
    assert not is_scanned("jd.docx", "")
    assert not is_scanned("jd.pdf", "x" * 200)


def test_text_is_wrapped_in_the_callers_tag() -> None:
    block = source_block(b"PK", "jd.docx", "Kế toán trưởng", "jd")

    assert block == {"type": "text", "text": "<jd>\nKế toán trưởng\n</jd>"}


def test_a_pdf_goes_as_a_document_without_leading_junk() -> None:
    block = source_block(b"\r\n%PDF-1.4 x", "jd.pdf", "", "jd")

    assert block is not None
    assert block["type"] == "document"
    assert base64.standard_b64decode(block["source"]["data"]) == b"%PDF-1.4 x"  # type: ignore[typeddict-item]


def test_no_text_and_no_pdf_means_nothing_to_send() -> None:
    assert source_block(b"\xd0\xcf", "jd.doc", "", "jd") is None


def test_a_reply_returns_fields_model_and_tokens() -> None:
    read = ask(reply())

    assert read.fields == {"title": "Kế toán trưởng"}
    assert read.model == "claude-sonnet-5-5-20261001"
    assert (read.input_tokens, read.output_tokens) == (900, 120)


def test_an_unusable_reply_keeps_its_cost_and_logs_the_label(
    caplog: pytest.LogCaptureFixture,
) -> None:
    read = ask(reply(stop_reason="max_tokens"))

    assert read.fields is None
    assert read.output_tokens == 120
    assert "Claude JD parse stopped with max_tokens" in caplog.text


def test_an_api_error_returns_nothing() -> None:
    assert ask(error=anthropic.APIConnectionError(request=None)) is None


def test_a_broken_reply_is_logged_without_its_content(caplog: pytest.LogCaptureFixture) -> None:
    assert ask(error=ValueError("raw reply: Velapay secret salary 99")) is None
    assert "Claude JD parse returned unusable output (ValueError)" in caplog.text
    assert "Velapay" not in caplog.text


def docx_upload(build: Any) -> SimpleUploadedFile:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return SimpleUploadedFile("jd.docx", buffer.getvalue())


def test_a_docx_reads_in_order_with_its_tables_in_place() -> None:
    # Tables used to come after every paragraph, far from the heading they belong to.
    def build(document: Any) -> None:
        document.add_paragraph("Mô tả công việc")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Chức danh"
        table.cell(0, 1).text = "Kế toán trưởng"
        document.add_paragraph("Yêu cầu")

    lines = extract_text(docx_upload(build)).splitlines()

    assert lines == ["Mô tả công việc", "Chức danh | Kế toán trưởng", "Yêu cầu"]


def test_a_merged_cell_is_read_once() -> None:
    # A JD form merges its summary across the row; it came back three times.
    def build(document: Any) -> None:
        table = document.add_table(rows=3, cols=3)
        summary = table.cell(0, 0).merge(table.cell(0, 2))
        summary.text = "Tóm lược công việc"
        label = table.cell(1, 0).merge(table.cell(2, 0))
        label.text = "Nhiệm vụ"
        table.cell(1, 1).text = "Lập báo cáo tài chính"
        table.cell(2, 1).text = "Quản lý dòng tiền"

    text = extract_text(docx_upload(build))

    assert text.count("Tóm lược công việc") == 1
    assert text.count("Nhiệm vụ") == 1
    assert "Nhiệm vụ | Lập báo cáo tài chính" in text
    assert "Quản lý dòng tiền" in text


def test_a_short_two_line_label_cell_stays_on_one_line() -> None:
    def build(document: Any) -> None:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Tên chức danh\nPosition Title"
        table.cell(0, 1).text = "Trưởng phòng QA"

    assert extract_text(docx_upload(build)) == "Tên chức danh / Position Title | Trưởng phòng QA"


def test_list_paragraphs_keep_a_bullet() -> None:
    def build(document: Any) -> None:
        document.add_paragraph("Quyền lợi")
        document.add_paragraph("Thưởng tháng 13", style="List Bullet")
        document.add_paragraph("• Bảo hiểm đầy đủ", style="List Bullet")

    lines = extract_text(docx_upload(build)).splitlines()

    assert lines == ["Quyền lợi", "- Thưởng tháng 13", "• Bảo hiểm đầy đủ"]
